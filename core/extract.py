"""Извлечение текста и метаданных из файла.

Никаких внешних сервисов: только локальные библиотеки. Если текст извлечь
нельзя (скан, картинка, старый .doc), файл помечается как требующий глазной
проверки или зрения модели — угадывать по имени регламент запрещает.
"""
from __future__ import annotations

import codecs
import html
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".log", ".ini", ".json", ".xml", ".html", ".htm"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic"}
LEGACY_EXT = {".doc", ".wps", ".pages"}           # .rtf и .odt читаем сами
MEDIA_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"}
OFFICE_EXT = {".docx", ".xlsx", ".xlsm", ".pptx", ".docm"}
IGNORED_PREFIX = re.compile(r"^(~\$|._)")

YEAR_RE = re.compile(r"(?<!\d)(19[89]\d|20[0-4]\d)(?!\d)")
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[.\-/ ](\d{1,2})[.\-/ ](\d{4})(?!\d)")
ISO_RE = re.compile(r"(20[0-4]\d|19[89]\d)-(\d{2})-(\d{2})")
RU_DATE_RE = re.compile(
    r"(?<!\d)(\d{1,2})[\s\-.]*(январ\w*|феврал\w*|март\w*|апрел\w*|ма[йя]|июн\w*|июл\w*|август\w*"
    r"|сентябр\w*|октябр\w*|ноябр\w*|декабр\w*)[\s,]*(\d{4})",
    re.IGNORECASE,
)
MONTH_STEMS = (
    ("январ", 1), ("феврал", 2), ("март", 3), ("апрел", 4), ("ма", 5), ("июн", 6),
    ("июл", 7), ("август", 8), ("сентябр", 9), ("октябр", 10), ("ноябр", 11), ("декабр", 12),
)


def month_number(word: str) -> int:
    low = word.lower()
    for stem, number in MONTH_STEMS:
        if low.startswith(stem):
            return number
    return 0


@dataclass
class Extracted:
    text: str = ""
    kind: str = "other"          # text | office | pdf | pdf-scan | image | legacy | empty | error
    needs_vision: bool = False
    force_review: bool = False
    note: str = ""
    years: list[str] = field(default_factory=list)
    dates: list[str] = field(default_factory=list)
    pages: int = 0
    meta: dict = field(default_factory=dict)

    @property
    def usable(self) -> bool:
        return bool(self.text.strip())


def clean_ws(value: str) -> str:
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t\u00a0]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _iso(year: str, month: int, day: int) -> str:
    if not 1 <= month <= 12 or not 1 <= day <= 31:
        return ""
    return f"{int(year):04d}-{month:02d}-{day:02d}"


def pick_dates(text: str) -> tuple[list[str], list[str]]:
    years = sorted(set(YEAR_RE.findall(text)))
    dates: list[str] = []
    for day, month, year in DATE_RE.findall(text):
        iso = _iso(year, int(month), int(day))
        if iso:
            dates.append(iso)
    for day, month_word, year in RU_DATE_RE.findall(text):
        iso = _iso(year, month_number(month_word), int(day))
        if iso:
            dates.append(iso)
    for year, month, day in ISO_RE.findall(text):
        iso = _iso(year, int(month), int(day))
        if iso:
            dates.append(iso)
    return years, sorted(set(dates))[:16]


# частые русские куски: в правильной кодировке их много, в неправильной почти нет
_COMMON = (" и ", " в ", " на ", " не ", " по ", " что ", " для ", " от ", " с ", " к ", "ени", "ого",
           "ост", "ств", "ние", "ова", "ать", "ами", "ото", "ают", "ной", "ому", "ний")


def _legacy_score(text: str) -> float:
    """Насколько текст похож на русский: частые слова вверх, «мусорные» буквы и псевдографика вниз."""
    low = text.casefold()
    common = sum(low.count(w) for w in _COMMON)
    junk = 0
    for ch in text:
        code = ord(ch)
        if 0x2500 <= code <= 0x259F:                       # псевдографика: так выглядит cp866 в чужой кодировке
            junk += 1
        elif code > 127 and ch.isalpha() and not (0x410 <= code <= 0x44F or ch in "Ёё"):
            junk += 1                                      # Ђ, Ґ, ў, Ï… — в нормальном русском тексте их нет
    cyr_lower = sum(1 for ch in low if "а" <= ch <= "я")
    return common * 3 + cyr_lower * 0.05 - junk * 6


def _utf16_order(head: bytes) -> str:
    """UTF-16 без BOM: латиница даёт нулевые байты, кириллица — байт 0x04 (старший байт U+04xx) через один."""
    if len(head) < 8:
        return ""
    evens, odds = head[0::2], head[1::2]
    if head.count(b"\x00") > len(head) * 0.25:
        return "utf-16-le" if odds.count(0) > evens.count(0) else "utf-16-be"
    if odds.count(4) > len(odds) * 0.3:
        return "utf-16-le"
    if evens.count(4) > len(evens) * 0.3:
        return "utf-16-be"
    return ""


def decode_text(raw: bytes) -> str:
    """Байты текстового файла → строка. Понимает BOM, UTF-16 без BOM, обрезанный UTF-8 и русские 8-битные.

    Файл мы читаем КУСКОМ (первые N КБ), поэтому хвост мог оборваться посреди символа — это не повод
    считать файл не-UTF-8 и читать его как cp1251 (раньше так и было: получался мусор)."""
    if raw.startswith(codecs.BOM_UTF8):
        return raw[3:].decode("utf-8", errors="ignore")
    if raw.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return raw[: len(raw) // 4 * 4].decode("utf-32", errors="ignore")
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return raw[: len(raw) // 2 * 2].decode("utf-16", errors="ignore")
    order = _utf16_order(raw[:4000])
    if order:
        return raw[: len(raw) // 2 * 2].decode(order, errors="ignore")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        if exc.end >= len(raw) - 3 and exc.start > 0:
            return raw[: exc.start].decode("utf-8", errors="ignore")      # оборван последний символ
    best_text, best_score = "", None
    for enc in ("cp1251", "cp866", "koi8-r"):             # cp1251 первой: ничья — её выигрыш
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        score = _legacy_score(text)
        if best_score is None or score > best_score * 1.25 + 2:
            best_text, best_score = text, score
    return best_text or raw.decode("utf-8", errors="replace")


def read_text_file(path: Path, limit: int) -> str:
    with path.open("rb") as stream:
        raw = stream.read(limit * 8)
    return decode_text(raw)


def docx_text(path: Path) -> str:
    """Текст docx: основной документ, колонтитулы (там номера и даты), без служебных полей и удалённого."""
    parts = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        members = ["word/document.xml"] + sorted(
            n for n in names if re.match(r"word/(header|footer)\d*\.xml$", n))
        for member in members:
            if member not in names:
                continue
            xml = archive.read(member).decode("utf-8", errors="replace")
            xml = re.sub(r"<w:instrText[^>]*>.*?</w:instrText>", "", xml, flags=re.S)   # PAGE, HYPERLINK…
            xml = re.sub(r"<w:del\b[^>]*>.*?</w:del>", "", xml, flags=re.S)           # правки «удалено»
            xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
            xml = re.sub(r"<w:br[^>]*/>|<w:cr/>", "\n", xml)
            xml = re.sub(r"</w:p>|</w:tr>", "\n", xml)
            xml = re.sub(r"</w:tc>", "\t", xml)
            xml = re.sub(r"<[^>]+>", "", xml)
            parts.append(html.unescape(xml))
    return "\n".join(parts)


_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "themedata", "colorschememapping",
             "latentstyles", "datastore", "listtable", "listoverridetable", "rsidtbl", "generator",
             "fldinst", "xmlnstbl", "mmathPr", "header", "footer", "headerl", "headerr", "footerl", "footerr"}


def rtf_text(path: Path, limit: int = 3_000_000) -> str:
    """Текст RTF без внешних библиотек: \\'hh (однобайтовая кодировка), \\uN, абзацы, служебные группы."""
    src = path.read_bytes()[:limit].decode("latin-1")
    cp_match = re.search(r"\\ansicpg(\d+)", src[:2000])
    codepage = f"cp{cp_match.group(1)}" if cp_match else "cp1251"
    out: list[str] = []
    stack: list[tuple[bool, int]] = []
    ignorable, uc, skip = False, 1, 0
    i, n = 0, len(src)
    word_re = re.compile(r"([A-Za-z]+)(-?\d+)? ?")
    while i < n:
        ch = src[i]
        if ch == "{":
            stack.append((ignorable, uc))
            i += 1
        elif ch == "}":
            ignorable, uc = stack.pop() if stack else (False, 1)
            i += 1
        elif ch == "\\":
            i += 1
            if i >= n:
                break
            nxt = src[i]
            if nxt in "\\{}":
                if not ignorable and skip == 0:
                    out.append(nxt)
                elif skip:
                    skip -= 1
                i += 1
            elif nxt == "'":
                code = src[i + 1:i + 3]
                i += 3
                if skip:
                    skip -= 1
                elif not ignorable:
                    try:
                        out.append(bytes([int(code, 16)]).decode(codepage, errors="replace"))
                    except ValueError:
                        pass
            elif nxt == "*":
                ignorable = True
                i += 1
            elif nxt.isalpha():
                match = word_re.match(src, i)
                word, num = match.group(1), match.group(2)
                i = match.end()
                if word in ("par", "line", "row", "sect", "page"):
                    out.append("\n") if not ignorable else None
                elif word == "tab":
                    out.append("\t") if not ignorable else None
                elif word == "cell":
                    out.append("\t") if not ignorable else None
                elif word == "u" and num is not None:
                    if not ignorable:
                        code = int(num)
                        out.append(chr(code + 65536 if code < 0 else code))
                    skip = uc
                elif word == "uc" and num is not None:
                    uc = int(num)
                elif word in _RTF_SKIP:
                    ignorable = True
            else:
                i += 1                                        # \~ \- \_ и прочие знаки
        elif ch in "\r\n":
            i += 1
        else:
            if skip:
                skip -= 1
            elif not ignorable:
                out.append(ch)
            i += 1
    return clean_ws("".join(out))


def odt_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("content.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"<text:tab/>", "\t", xml)
    xml = re.sub(r"<text:line-break/>|</text:p>|</text:h>|</table:table-row>", "\n", xml)
    xml = re.sub(r"</table:table-cell>", "\t", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def xlsx_text(path: Path, max_rows: int = 120) -> str:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""
    out = []
    try:
        for sheet in wb.worksheets[:6]:
            out.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    out.append("\t".join(cells))
                if len(out) > max_rows:
                    break
            if len(out) > max_rows:
                break
    finally:
        wb.close()
    return "\n".join(out)


def pptx_text(path: Path) -> str:
    out = []
    try:
        with zipfile.ZipFile(path) as archive:
            slides = sorted(
                (n for n in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                key=lambda n: int(re.search(r"(\d+)", n).group(1)),
            )
            for member in slides[:20]:
                xml = archive.read(member).decode("utf-8", errors="replace")
                xml = re.sub(r"</a:p>", "\n", xml)
                out.append(html.unescape(re.sub(r"<[^>]+>", "", xml)))
    except Exception:
        return ""
    return "\n".join(out)


def pdf_text(path: Path, max_pages: int = 12) -> tuple[str, int, bool]:
    """Возвращает (текст, страниц, это_скан)."""
    try:
        import fitz
    except Exception:
        return "", 0, False
    try:
        with fitz.open(path) as doc:
            pages = doc.page_count
            chunk = []
            for i in range(min(pages, max_pages)):
                chunk.append(doc[i].get_text())
            text = "\n".join(chunk)
            images = sum(len(doc[i].get_images(full=True)) for i in range(min(pages, max_pages)))
        scan = len(text.strip()) < 250 and (pages == 0 or images >= pages)
        return text, pages, scan
    except Exception:
        return "", 0, False


def image_meta(path: Path) -> dict:
    meta: dict = {}
    try:
        from PIL import ExifTags, Image

        with Image.open(path) as img:
            meta["width"], meta["height"] = img.size
            meta["format"] = img.format or path.suffix.lstrip(".").upper()
            exif = img.getexif()
            for tag, value in exif.items():
                name = ExifTags.TAGS.get(tag, str(tag))
                if name in ("DateTime", "DateTimeOriginal", "Model", "Make", "ImageDescription") and value:
                    meta[name] = str(value).strip("\x00 ")[:120]
    except Exception:
        pass
    return meta


def head_tail(text: str, limit: int) -> str:
    """Начало важнее, но дата/подпись часто в конце — берём и там, и там."""
    text = clean_ws(text)
    if len(text) <= limit:
        return text
    head = int(limit * 0.72)
    tail = int(limit * 0.2)
    return text[:head] + "\n…[середина пропущена]…\n" + text[-tail:]


DOC_HEADWORDS = (
    "РАПОРТ", "ОТЧЕТ", "ОТЧЁТ", "ПРОШЕНИЕ", "АКТ", "СМЕТ", "ДОГОВОР", "ЗАЯВКА", "ПИСЬМО",
    "ПРИГЛАШЕНИЕ", "БЛАГОДАР", "ПРОТОКОЛ", "ПОЛОЖЕНИЕ", "СПИСОК", "СЦЕНАРИЙ", "ХОДАТАЙСТВ",
    "СОГЛАСИЕ", "ИНФОРМ", "ЦИРКУЛ", "ПЛАН", "СПРАВКА",
)


def multi_doc_markers(text: str) -> list[str]:
    """Признаки того, что в одном файле подшиты несколько документов."""
    markers: list[str] = []
    if not text:
        return markers
    heads = sum(
        1 for line in text.splitlines()
        if 3 <= len(line.strip()) <= 40 and any(line.strip().upper().startswith(word) for word in DOC_HEADWORDS)
    )
    if heads >= 2:
        markers.append(f"заголовков документов: {heads}")
    pages = text.count("\f")
    if pages >= 2:
        markers.append(f"переводов страниц: {pages}")
    if len(re.findall(r"(?i)приложение\s*[:№]", text)) >= 2:
        markers.append("несколько приложений")
    dates = set(re.findall(r"(?:19|20)\d{2}-\d{2}-\d{2}", text))
    if len(dates) >= 3:
        markers.append(f"разных дат: {len(dates)}")
    return markers


def extract(path: Path, content_chars: int = 3200, max_pdf_pages: int = 12,
            ocr: bool = False, media_meta: bool = False) -> Extracted:
    """Прочитать файл и вернуть текст и признаки, достаточные для авто-решения."""
    res = Extracted()
    ext = path.suffix.lower()
    if not path.exists():
        res.kind, res.force_review, res.note = "error", True, "файл не найден на диске"
        return res
    if path.stat().st_size == 0:
        res.kind, res.force_review, res.note = "empty", True, "пустой файл"
        return res

    try:
        if ext in TEXT_EXT:
            res.kind = "text"
            res.text = read_text_file(path, content_chars)
        elif ext in {".docx", ".docm"}:
            res.kind = "office"
            res.text = docx_text(path)
        elif ext == ".rtf":
            res.kind = "office"
            res.text = rtf_text(path)
            if len(res.text.strip()) < 15:                    # не RTF по сути (пустой/картинка)
                res.kind, res.force_review = "legacy", True
                res.note = "RTF без читаемого текста — сначала конвертировать в docx/pdf"
        elif ext == ".odt":
            res.kind = "office"
            res.text = odt_text(path)
        elif ext in {".xlsx", ".xlsm"}:
            res.kind = "office"
            res.text = xlsx_text(path)
        elif ext == ".pptx":
            res.kind = "office"
            res.text = pptx_text(path)
        elif ext == ".pdf":
            text, pages, scan = pdf_text(path, max_pdf_pages)
            res.pages = pages
            res.kind = "pdf-scan" if scan else "pdf"
            res.text = text
            res.needs_vision = scan
            if scan:
                res.note = "PDF без текстового слоя (скан) — нужен просмотр или OCR"
        elif ext in IMAGE_EXT:
            res.kind = "image"
            if ext in (".heic", ".heif"):
                from .plugins import ensure_heif

                ensure_heif()
            res.meta = image_meta(path)
            res.needs_vision = True
            stamp = res.meta.get("DateTimeOriginal") or res.meta.get("DateTime") or ""
            res.text = f"Изображение {res.meta.get('format', ext)} {res.meta.get('width', '?')}x{res.meta.get('height', '?')}. EXIF дата: {stamp or 'нет'}"
            res.note = "фото/картинка"
        elif ext in LEGACY_EXT:
            res.kind = "legacy"
            res.force_review = True
            res.note = f"устаревший формат {ext} — сначала конвертировать в docx/pdf"
        elif ext in MEDIA_EXT:
            res.kind = "media"
            res.needs_vision = False
            if media_meta:
                from .plugins import media_info

                res.meta = media_info(path)
            stamp = res.meta.get("название") or ""
            described = " ".join(f"{k}: {v}" for k, v in res.meta.items())
            res.text = f"Медиафайл {ext}" + (f". {described}" if described else "")
            res.note = "аудио/видео — содержимое не читается" + (f"; название из тегов: {stamp}" if stamp else "")
            res.force_review = True
        else:
            res.kind = "other"
            res.force_review = True
            res.note = f"нераспознанный тип {ext or 'без расширения'}"
    except zipfile.BadZipFile:
        res.kind, res.force_review, res.note = "error", True, "повреждённый контейнер office/pdf"
    except Exception as exc:  # не роняем партию из-за одного файла
        res.kind, res.force_review, res.note = "error", True, f"не прочитан: {type(exc).__name__}: {exc}"

    if ocr and res.kind in ("pdf-scan", "image") and (not res.usable or res.kind == "pdf-scan"):
        from .plugins import ocr_text

        recognized = clean_ws(ocr_text(path) or "")
        if len(recognized) >= 40:
            res.text = recognized
            res.meta["ocr"] = True
            res.force_review = False
            res.needs_vision = False   # текст уже есть — «смотреть глазами» не требуется
            res.note = "текст получен распознаванием — проверьте ключевые данные"
            if res.kind == "pdf-scan":
                res.kind = "pdf-ocr"

    years, dates = pick_dates(res.text + " " + path.name)
    res.years, res.dates = years, dates
    multi = multi_doc_markers(res.text)
    if multi:
        res.meta["несколько документов"] = "; ".join(multi)
        res.note = (res.note + "; " if res.note else "") + "в файле несколько документов — проверьте целиком"
    if res.usable and res.kind != "image":
        res.text = head_tail(res.text, content_chars)
    if not res.usable and res.kind not in ("image", "error", "empty", "legacy"):
        res.force_review = True
        res.note = res.note or "текст не извлечён — нужен глаз"
    return res


def modified_year(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y")
